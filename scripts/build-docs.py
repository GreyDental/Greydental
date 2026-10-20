#!/usr/bin/env python3
"""Generate the Grey Dental architecture diagrams and Word document.

Usage:
    python scripts/build-docs.py
    npm run docs:toc        # optional: populate the Word TOC field

Inputs:
    docs/TECHNICAL_DOCUMENTATION.md   (source of truth)

Outputs:
    docs/diagrams/*.png               (architecture diagrams)
    docs/Grey_Dental_Technical_Documentation.docx

Requires: pillow, python-docx
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
MARKDOWN = DOCS / "TECHNICAL_DOCUMENTATION.md"
DOCX = DOCS / "Grey_Dental_Technical_Documentation.docx"

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------

INK = "#1F3B47"
MUTED = "#5B6B73"
ACCENT = "#3A738D"
ACCENT_DARK = "#2F5F75"
GREEN = "#3F9E7C"
GREEN_SOFT = "#E5F8F0"
BLUE_SOFT = "#E8F1F5"
AMBER = "#B7791F"
AMBER_SOFT = "#FFF6E5"
BORDER = "#C9D6DC"
PANEL = "#F8FBFC"

FONT_DIR = Path("C:/Windows/Fonts")
FONT_FILES = {
    (False, False): ["segoeui.ttf", "DejaVuSans.ttf", "arial.ttf"],
    (True, False): ["segoeuib.ttf", "DejaVuSans-Bold.ttf", "arialbd.ttf"],
    (False, True): ["segoeuii.ttf", "DejaVuSans-Oblique.ttf", "ariali.ttf"],
}


def _font_path(bold: bool, italic: bool) -> Path:
    candidates = FONT_FILES.get((bold, italic)) or FONT_FILES[(bold, False)]
    for name in candidates:
        path = FONT_DIR / name
        if path.exists():
            return path
    return FONT_DIR / candidates[0]


class Canvas:
    """Small drawing helper that works in logical pixels and renders at 2x."""

    SCALE = 2

    def __init__(self, width: int, height: int, background: str = "#FFFFFF"):
        self.width = width
        self.height = height
        self.image = Image.new(
            "RGB",
            (width * self.SCALE, height * self.SCALE),
            background,
        )
        self.draw = ImageDraw.Draw(self.image)
        self._fonts: dict[tuple[int, bool, bool], ImageFont.FreeTypeFont] = {}

    def font(self, size: int, bold: bool = False, italic: bool = False):
        key = (size, bold, italic)
        if key not in self._fonts:
            self._fonts[key] = ImageFont.truetype(
                str(_font_path(bold, italic)),
                size * self.SCALE,
            )
        return self._fonts[key]

    def _s(self, value: float) -> int:
        return int(round(value * self.SCALE))

    def text(
        self,
        x: float,
        y: float,
        value: str,
        size: int = 13,
        bold: bool = False,
        italic: bool = False,
        color: str = INK,
        anchor: str = "la",
    ) -> None:
        self.draw.text(
            (self._s(x), self._s(y)),
            value,
            font=self.font(size, bold, italic),
            fill=color,
            anchor=anchor,
        )

    def wrap(self, value: str, size: int, max_width: float) -> list[str]:
        font = self.font(size)
        words = value.split()
        if not words:
            return [""]
        lines: list[str] = []
        current = words[0]
        for word in words[1:]:
            trial = f"{current} {word}"
            if self.draw.textlength(trial, font=font) <= max_width * self.SCALE:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
        return lines

    def box(
        self,
        x: float,
        y: float,
        width: float,
        height: float,
        title: str,
        lines: list[str] | None = None,
        *,
        fill: str = "#FFFFFF",
        stroke: str = BORDER,
        title_color: str = ACCENT_DARK,
        body_color: str = MUTED,
        radius: int = 14,
        border_width: int = 2,
        title_size: int = 15,
        body_size: int = 12,
        padding: int = 18,
        center: bool = False,
    ) -> None:
        self.draw.rounded_rectangle(
            [
                self._s(x),
                self._s(y),
                self._s(x + width),
                self._s(y + height),
            ],
            radius=self._s(radius),
            fill=fill,
            outline=stroke,
            width=self._s(border_width),
        )
        text_x = x + width / 2 if center else x + padding
        anchor = "ma" if center else "la"
        self.text(
            text_x,
            y + padding,
            title,
            size=title_size,
            bold=True,
            color=title_color,
            anchor=anchor,
        )
        cursor = y + padding + title_size + 10
        for line in lines or []:
            wrapped = self.wrap(line, body_size, width - padding * 2)
            for part in wrapped:
                self.text(
                    text_x,
                    cursor,
                    part,
                    size=body_size,
                    color=body_color,
                    anchor=anchor,
                )
                cursor += body_size + 6

    def diamond(
        self,
        cx: float,
        cy: float,
        width: float,
        height: float,
        title: str,
        *,
        fill: str = AMBER_SOFT,
        stroke: str = AMBER,
        color: str = AMBER,
        size: int = 14,
    ) -> None:
        points = [
            (self._s(cx), self._s(cy - height / 2)),
            (self._s(cx + width / 2), self._s(cy)),
            (self._s(cx), self._s(cy + height / 2)),
            (self._s(cx - width / 2), self._s(cy)),
        ]
        self.draw.polygon(points, fill=fill, outline=stroke)
        self.draw.line(points + [points[0]], fill=stroke, width=self._s(2))
        self.text(cx, cy, title, size=size, bold=True, color=color, anchor="mm")

    def _arrow_head(
        self,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        color: str,
        length: float = 12,
    ) -> None:
        angle = math.atan2(y2 - y1, x2 - x1)
        left = (
            x2 - length * math.cos(angle - math.pi / 7),
            y2 - length * math.sin(angle - math.pi / 7),
        )
        right = (
            x2 - length * math.cos(angle + math.pi / 7),
            y2 - length * math.sin(angle + math.pi / 7),
        )
        self.draw.polygon(
            [
                (self._s(x2), self._s(y2)),
                (self._s(left[0]), self._s(left[1])),
                (self._s(right[0]), self._s(right[1])),
            ],
            fill=color,
        )

    def arrow(
        self,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        *,
        color: str = ACCENT,
        width: int = 2,
        dashed: bool = False,
        head: bool = True,
    ) -> None:
        if dashed:
            total = math.hypot(x2 - x1, y2 - y1)
            step = 10
            drawn = 0
            while drawn < total:
                start = drawn / total
                end = min((drawn + 6) / total, 1)
                self.draw.line(
                    [
                        self._s(x1 + (x2 - x1) * start),
                        self._s(y1 + (y2 - y1) * start),
                        self._s(x1 + (x2 - x1) * end),
                        self._s(y1 + (y2 - y1) * end),
                    ],
                    fill=color,
                    width=self._s(width),
                )
                drawn += step
        else:
            self.draw.line(
                [self._s(x1), self._s(y1), self._s(x2), self._s(y2)],
                fill=color,
                width=self._s(width),
            )
        if head:
            self._arrow_head(x1, y1, x2, y2, color)

    def label(
        self,
        x: float,
        y: float,
        value: str,
        *,
        size: int = 12,
        color: str = MUTED,
        background: str = "#FFFFFF",
        anchor: str = "mm",
        bold: bool = False,
    ) -> None:
        font = self.font(size, bold)
        bbox = self.draw.textbbox((0, 0), value, font=font, anchor=anchor)
        pad = self._s(4)
        self.draw.rectangle(
            [
                self._s(x) + bbox[0] - pad,
                self._s(y) + bbox[1] - pad,
                self._s(x) + bbox[2] + pad,
                self._s(y) + bbox[3] + pad,
            ],
            fill=background,
        )
        self.text(x, y, value, size=size, color=color, anchor=anchor, bold=bold)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.image.save(path, "PNG", optimize=True)


# ---------------------------------------------------------------------------
# Diagrams
# ---------------------------------------------------------------------------


def system_context(path: Path) -> None:
    c = Canvas(1400, 800)

    c.box(
        60,
        150,
        260,
        160,
        "Student / admin browser",
        ["Marketing site", "Student dashboard", "Admin dashboard"],
        fill=BLUE_SOFT,
        title_color=ACCENT_DARK,
    )

    c.draw.rounded_rectangle(
        [c._s(380), c._s(80), c._s(1010), c._s(700)],
        radius=c._s(20),
        outline=ACCENT,
        width=c._s(2),
    )
    c.text(400, 95, "Next.js 16 application · Vercel", size=14, bold=True, color=ACCENT)

    c.box(
        410,
        140,
        580,
        120,
        "Pages & React components",
        [
            "App Router · server and client components",
            "Tailwind CSS v4 · framer-motion · lucide-react",
        ],
        fill="#FFFFFF",
    )
    c.box(
        410,
        300,
        580,
        105,
        "Middleware",
        [
            "Refreshes the Supabase session cookie on every request",
            "Guards /student-dashboard, /admin-dashboard and /checkout",
        ],
        fill="#FFFFFF",
    )
    c.box(
        410,
        445,
        580,
        145,
        "API routes · /api/*",
        [
            "checkout · courses · enrollments · workshops · resources",
            "learners · assessments · profile · auth/register",
            "webhooks/stripe",
        ],
        fill="#FFFFFF",
    )

    c.box(
        1080,
        130,
        280,
        280,
        "Supabase",
        [
            "Auth — email/password + Google OAuth",
            "Postgres — tables, constraints, RLS policies",
            "Storage — uploaded resources",
        ],
        fill=GREEN_SOFT,
        stroke=GREEN,
        title_color=GREEN,
    )
    c.box(
        1080,
        470,
        280,
        140,
        "Stripe",
        ["Checkout Sessions", "Signed webhook events"],
        fill=AMBER_SOFT,
        stroke=AMBER,
        title_color=AMBER,
    )

    c.arrow(320, 230, 406, 230)
    c.label(363, 210, "HTTPS")

    c.arrow(700, 260, 700, 296)
    c.label(745, 278, "client fetch")

    c.arrow(700, 405, 700, 441)
    c.label(748, 423, "route handlers")

    c.arrow(992, 500, 1076, 390)
    c.label(1045, 425, "SQL via @supabase/ssr")

    c.arrow(992, 545, 1076, 520)

    c.arrow(1076, 560, 992, 585, dashed=True)

    # Browser -> Supabase Auth (direct), routed above the app container.
    c.arrow(120, 150, 120, 40, head=False)
    c.arrow(120, 40, 1210, 40, head=False)
    c.arrow(1210, 40, 1210, 126)
    c.label(660, 24, "sign-up, login and OAuth (direct browser → Supabase Auth)")

    c.text(
        60,
        735,
        "All non-auth data access goes through /api/*. The browser only talks to Supabase directly for authentication.",
        size=12,
        italic=True,
        color=MUTED,
    )
    c.save(path)


def request_lifecycle(path: Path) -> None:
    c = Canvas(1400, 800)

    lifelines = [
        (170, "Browser", 60, 220),
        (450, "Middleware", 340, 220),
        (730, "API route", 620, 220),
        (1010, "Supabase Auth", 900, 220),
        (1275, "Postgres + RLS", 1180, 190),
    ]

    for x, title, left, width in lifelines:
        c.box(
            left,
            60,
            width,
            72,
            title,
            None,
            fill=BLUE_SOFT,
            center=True,
            title_size=14,
            padding=14,
        )
        c.arrow(x, 132, x, 720, dashed=True, head=False, color=BORDER)

    x_browser, x_mw, x_api, x_auth, x_db = (x for x, *_ in lifelines)

    steps = [
        (180, x_browser, x_mw, "1. Request a protected page", False),
        (250, x_mw, x_auth, "2. getUser() refreshes the session", False),
        (320, x_mw, x_browser, "3. No session → 302 /login?next=…", True),
        (395, x_browser, x_api, "4. GET /api/courses with the session cookie", False),
        (465, x_api, x_auth, "5. Read cookies via @supabase/ssr", False),
        (535, x_api, x_db, "6. Query the requested tables", False),
        (605, x_db, x_api, "7. Return only rows allowed by RLS", True),
        (675, x_api, x_browser, "8. 200 JSON response", True),
    ]

    for y, x1, x2, label, returning in steps:
        c.arrow(x1, y, x2, y, dashed=returning, color=ACCENT if not returning else MUTED)
        c.label((x1 + x2) / 2, y - 17, label, size=12)

    c.text(
        60,
        760,
        "If NEXT_PUBLIC_SUPABASE_URL or the anon key is missing, middleware fails closed on protected paths.",
        size=12,
        italic=True,
        color=MUTED,
    )
    c.save(path)


def purchase_flow(path: Path) -> None:
    c = Canvas(1400, 900)

    c.box(
        70,
        50,
        330,
        80,
        "Checkout page",
        ["/checkout?course=<slug>"],
        fill=BLUE_SOFT,
        title_size=15,
    )
    c.box(
        70,
        190,
        330,
        90,
        "POST /api/checkout",
        ["Validates the session and course slug"],
        fill="#FFFFFF",
    )
    c.arrow(235, 130, 235, 186)
    c.label(300, 158, "submit")

    c.diamond(640, 235, 300, 150, "Stripe demo mode?")

    c.arrow(400, 235, 490, 235)
    c.label(440, 215, "request")

    c.box(
        440,
        400,
        360,
        105,
        "orders.ts · completeEnrollmentPurchase()",
        ["Inserts the order and enrollment directly"],
        fill=GREEN_SOFT,
        stroke=GREEN,
        title_color=GREEN,
        title_size=14,
    )
    c.arrow(640, 310, 640, 396)
    c.label(690, 352, "yes (dummy Stripe key)")

    c.box(
        990,
        190,
        340,
        100,
        "Stripe Checkout Session",
        ["Hosted payment page"],
        fill=AMBER_SOFT,
        stroke=AMBER,
        title_color=AMBER,
    )
    c.arrow(790, 235, 986, 235)
    c.label(888, 215, "no (live keys)")

    c.box(
        990,
        360,
        340,
        100,
        "POST /api/webhooks/stripe",
        ["checkout.session.completed"],
        fill=AMBER_SOFT,
        stroke=AMBER,
        title_color=AMBER,
    )
    c.arrow(1160, 290, 1160, 356)
    c.label(1232, 322, "user pays", color=AMBER)

    c.box(
        600,
        585,
        430,
        105,
        "orders.status = 'paid'",
        ["One active enrollments row per learner and course"],
        fill=GREEN_SOFT,
        stroke=GREEN,
        title_color=GREEN,
    )
    c.arrow(620, 505, 680, 581)
    c.arrow(1160, 460, 980, 581)
    c.label(1070, 540, "verify signature", color=AMBER)

    c.box(
        640,
        760,
        350,
        85,
        "Student sees the course",
        ["/student-dashboard/courses"],
        fill=BLUE_SOFT,
    )
    c.arrow(815, 690, 815, 756)

    c.text(
        70,
        860,
        "Both branches converge on the same helper, so demo and live purchases produce identical order and enrollment rows.",
        size=12,
        italic=True,
        color=MUTED,
    )
    c.save(path)


def data_model(path: Path) -> None:
    c = Canvas(1400, 900)

    c.box(
        540,
        50,
        300,
        115,
        "profiles",
        ["id (PK) · role · full_name", "institution · specialty · phone"],
        fill=BLUE_SOFT,
        center=True,
    )

    c.box(
        50,
        280,
        370,
        260,
        "Courses & content",
        [
            "courses — id (PK), slug, title,",
            "price_cents, module_count, status",
            "enrollments — user_id (FK), course_id (FK),",
            "progress, status",
            "resources — course_id (FK), file_url, is_public",
        ],
        fill=PANEL,
        title_color=ACCENT,
    )

    c.box(
        500,
        280,
        370,
        220,
        "Commerce & credentials",
        [
            "orders — user_id (FK), course_id (FK),",
            "amount_cents, status",
            "certificates — user_id (FK), course_id (FK),",
            "issued_at, certificate_url",
        ],
        fill=PANEL,
        title_color=ACCENT,
    )

    c.box(
        960,
        280,
        390,
        220,
        "Live workshops",
        [
            "workshops — id (PK), slug, starts_at,",
            "capacity, status",
            "workshop_registrations — workshop_id (FK),",
            "user_id (FK)",
        ],
        fill=PANEL,
        title_color=ACCENT,
    )

    c.box(
        960,
        590,
        390,
        240,
        "Assessments",
        [
            "assessments — course_id (FK), title, status",
            "assessment_questions — assessment_id (FK),",
            "prompt, correct_index",
            "assessment_submissions — assessment_id (FK),",
            "user_id (FK), score, cheated",
        ],
        fill=PANEL,
        title_color=ACCENT,
    )

    # profiles -> each domain
    c.arrow(590, 165, 300, 276)
    c.label(430, 205, "1:N", size=12)

    c.arrow(690, 165, 690, 276)
    c.label(726, 220, "1:N", size=12)

    c.arrow(790, 165, 1080, 276)
    c.label(975, 202, "1:N", size=12)

    # profiles -> assessment submissions, routed down the right margin
    c.arrow(840, 108, 1375, 108, head=False, color=MUTED, dashed=True)
    c.arrow(1375, 108, 1375, 745, head=False, color=MUTED, dashed=True)
    c.arrow(1375, 745, 1354, 745, color=MUTED, dashed=True)
    c.label(1260, 96, "submissions.user_id (1:N)", size=11, color=MUTED)

    # courses -> assessments
    c.arrow(235, 540, 235, 565, head=False, color=MUTED, dashed=True)
    c.arrow(235, 565, 900, 565, head=False, color=MUTED, dashed=True)
    c.arrow(900, 565, 900, 700, head=False, color=MUTED, dashed=True)
    c.arrow(900, 700, 956, 700, color=MUTED, dashed=True)
    c.label(570, 553, "assessments.course_id (N:1)", size=11, color=MUTED)

    c.text(
        50,
        860,
        "PK = primary key · FK = foreign key · 1:N = one-to-many · N:1 = many-to-one",
        size=12,
        italic=True,
        color=MUTED,
    )
    c.save(path)


def build_diagrams() -> None:
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    system_context(DIAGRAMS / "system-context.png")
    request_lifecycle(DIAGRAMS / "request-lifecycle.png")
    purchase_flow(DIAGRAMS / "purchase-flow.png")
    data_model(DIAGRAMS / "data-model.png")
    print(f"Wrote 4 diagrams to {DIAGRAMS}")


# ---------------------------------------------------------------------------
# Markdown -> DOCX
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT.replace("#", ""))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, value: str, *, size: float | None = None) -> None:
    for token in INLINE_RE.split(value):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x2F, 0x5F, 0x75)
        elif token.startswith("[") and "](" in token:
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if match:
                add_hyperlink(paragraph, match.group(1), match.group(2))
        else:
            paragraph.add_run(token)
        if size is not None:
            for run in paragraph.runs:
                run.font.size = Pt(size)


def shade(cell_or_paragraph, fill: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:color"), "auto")
    element.set(qn("w:fill"), fill.replace("#", ""))
    if hasattr(cell_or_paragraph, "_tc"):
        cell_or_paragraph._tc.get_or_add_tcPr().append(element)
    else:
        cell_or_paragraph._p.get_or_add_pPr().append(element)


def add_page_number_footer(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Grey Dental — Technical Documentation  ·  ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x79)

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x79)


def add_toc(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-2" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = (
        "Right-click here and choose Update Field to build the table of contents."
    )
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    for name, size, color, space_before in (
        ("Heading 1", 17, ACCENT_DARK, 18),
        ("Heading 2", 13.5, ACCENT, 14),
        ("Heading 3", 11.5, ACCENT_DARK, 10),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    # Word reuses these custom styles when it builds the TOC field. Keeping the
    # entries compact ensures the contents fit on one page.
    for name, size, indent in (("TOC 1", 9.5, 0), ("TOC 2", 9, 0.22)):
        try:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        except ValueError:
            style = document.styles[name]
        style.base_style = document.styles["Normal"]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(1)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.left_indent = Inches(indent)
        style.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.7),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )


def build_cover(document: Document) -> None:
    for _ in range(5):
        document.add_paragraph()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Grey Dental")
    run.font.name = "Calibri"
    run.font.size = Pt(42)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT_DARK.replace("#", ""))

    subtitle = document.add_paragraph()
    run = subtitle.add_run("Technical Documentation")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(ACCENT.replace("#", ""))

    document.add_paragraph()
    tagline = document.add_paragraph()
    run = tagline.add_run(
        "A standards-compliant technical reference for the Grey Dental platform"
    )
    run.font.size = Pt(12.5)
    run.italic = True
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x79)

    for _ in range(8):
        document.add_paragraph()

    for line in (
        "Prepared following the Intellema Technical Documentation Standard",
        "Version 1.1  ·  October 2026  ·  Internal",
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x77, 0x77, 0x79)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            value = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            add_inline(paragraph, value, size=9.5)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shade(cell, "E8F1F5")
    document.add_paragraph()


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    shade(paragraph, "F4F7F8")
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        if index < len(lines) - 1:
            run.add_break()


def build_body(document: Document) -> None:
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    first_heading = True

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            index += 1
            continue

        image = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            alt, source = image.groups()
            image_path = (DOCS / source).resolve()
            if image_path.suffix.lower() == ".svg":
                image_path = image_path.with_suffix(".png")
            if image_path.exists():
                picture = document.add_picture(str(image_path), width=Inches(6.5))
                picture_paragraph = document.paragraphs[-1]
                picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture_paragraph.paragraph_format.space_after = Pt(3)
                if alt:
                    caption = document.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = caption.add_run(alt)
                    run.italic = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x77, 0x77, 0x79)
            else:
                print(f"Warning: diagram not found: {image_path}")
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1:
                paragraph = document.add_heading(level=1)
                add_inline(paragraph, text)
                if first_heading:
                    paragraph.paragraph_format.page_break_before = True
                    first_heading = False
                for run in paragraph.runs:
                    run.font.size = Pt(22)
                    run.font.color.rgb = RGBColor.from_string(
                        ACCENT_DARK.replace("#", "")
                    )
            elif level == 2:
                paragraph = document.add_heading(level=1)
                add_inline(paragraph, text)
            else:
                paragraph = document.add_heading(level=2)
                add_inline(paragraph, text)
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue

        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.22)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(f"{ordered.group(1)}.  ")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(ACCENT.replace("#", ""))
            add_inline(paragraph, ordered.group(2))
            index += 1
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run("•  ")
            run.font.color.rgb = RGBColor.from_string(GREEN.replace("#", ""))
            add_inline(paragraph, stripped[2:])
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.08
        add_inline(paragraph, stripped)
        index += 1


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    configure_styles(document)
    add_page_number_footer(section)

    document.core_properties.title = "Grey Dental — Technical Documentation"
    document.core_properties.subject = "Architecture and implementation reference"
    document.core_properties.author = "Grey Dental engineering"
    document.core_properties.version = "1.1"

    build_cover(document)

    toc_heading = document.add_paragraph()
    toc_heading.paragraph_format.page_break_before = True
    run = toc_heading.add_run("Table of Contents")
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ACCENT_DARK.replace("#", ""))
    add_toc(document)

    build_body(document)

    document.save(DOCX)
    print(f"Wrote {DOCX}")


if __name__ == "__main__":
    build_diagrams()
    build_docx()
